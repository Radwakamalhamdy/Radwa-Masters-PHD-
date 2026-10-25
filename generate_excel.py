# -*- coding: utf-8 -*-
import sys
import io
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

with open('curated_scholarships.json', 'r', encoding='utf-8') as f:
    items = json.load(f)

print(f"Generating Comprehensive Excel file for {len(items)} scholarships...")

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "دليل المنح والمطابقة لرضوى"
ws.views.sheetView[0].rightToLeft = True

# Define styles
header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

category_colors = {
    "أفريقيا والدول النامية": "FEF3C7", # Light amber
    "المنح الحكومية والدولية الكبرى": "DBEAFE", # Light blue
    "مراكز ومعاهد التكنولوجيا الحيوية والعلوم الطبية": "DCFCE7", # Light green
    "منح الجامعات العالمية والمؤسسات المرموقة": "F3E8FF", # Light purple
    "محركات البحث وأدلة وظائف الدكتوراه": "FCE7F3", # Light pink
}

match_colors = {
    "98%": "D1FAE5", # Emerald light
    "96%": "D1FAE5",
    "95%": "D1FAE5",
    "94%": "DCFCE7",
    "92%": "E0F2FE", # Blue light
    "90%": "E0F2FE",
    "88%": "FEF3C7", # Amber light
    "85%": "FEF3C7",
    "80%": "FFEDD5"  # Orange light
}

thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

headers = [
    "م",
    "اسم المنحة / الجهة المانحة",
    "التصنيف والنطاق",
    "المرحلة الدراسية",
    "نسبة التوافق مع الـ CV",
    "حالة الاستيفاء والجاهزية",
    "المتطلبات التفصيلية والشروط الأكاديمية",
    "طرح الـ CV: المتبقي لرضوى لاستيفاء الشروط 100%",
    "الرابط الرسمي المباشر",
    "ميعاد بداية التقديم",
    "آخر ميعاد للتقديم (Deadline)",
    "نوع التمويل",
    "ملاحظات توجيهية"
]

ws.append(headers)

# Format header row
for col_idx in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

ws.row_dimensions[1].height = 36

for idx, item in enumerate(items, 1):
    row_data = [
        idx,
        item.get("name_ar") if item.get("name_ar") and item.get("name_ar") != item.get("name") else item.get("name"),
        item.get("category", "عام"),
        item.get("degree", "ماجستير ودكتوراه"),
        item.get("match_rate", "90%"),
        item.get("readiness_status", "مؤهلة للتقديم"),
        item.get("detailed_requirements") or item.get("requirements", ""),
        item.get("cv_gap_analysis", "راجع الـ CV للتطابق المباشر"),
        item.get("link", ""),
        item.get("start_date", ""),
        item.get("deadline", ""),
        item.get("funding", "تمويل كامل أو جزئي"),
        item.get("notes", "")
    ]
    ws.append(row_data)
    row_num = idx + 1
    ws.row_dimensions[row_num].height = 95 # ample space for multi-line requirements and gap analysis

    cat = item.get("category", "")
    bg_color = category_colors.get(cat, "FFFFFF")
    cat_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")

    match_r = item.get("match_rate", "90%")
    m_color = match_colors.get(match_r, "FFFFFF")
    match_fill = PatternFill(start_color=m_color, end_color=m_color, fill_type="solid")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=row_num, column=col_idx)
        cell.font = Font(name="Calibri", size=10)
        cell.border = thin_border
        
        # Alignments
        if col_idx in [1, 3, 4, 5, 6, 10, 11]:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        elif col_idx in [7, 8]:
            cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
        else:
            cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
            
        # Highlight category cell
        if col_idx == 3:
            cell.fill = cat_fill
            cell.font = Font(name="Calibri", size=10, bold=True)
            
        # Highlight match rate
        if col_idx == 5:
            cell.fill = match_fill
            cell.font = Font(name="Calibri", size=10, bold=True, color="065F46" if "9" in match_r else "1E3A8A")

# Adjust column widths
col_widths = {
    1: 6,   # م
    2: 36,  # الاسم
    3: 24,  # التصنيف
    4: 18,  # الدرجة
    5: 16,  # نسبة التوافق
    6: 28,  # حالة الاستيفاء
    7: 55,  # المتطلبات التفصيلية
    8: 60,  # تحليل الفجوة وطرح الـ CV
    9: 32,  # الرابط
    10: 22, # البداية
    11: 24, # الموعد النهائي
    12: 28, # التمويل
    13: 35  # ملاحظات
}

for col_idx, width in col_widths.items():
    col_letter = get_column_letter(col_idx)
    ws.column_dimensions[col_letter].width = width

excel_path = "Scholarships_Master_PhD_Radwa_Complete.xlsx"
wb.save(excel_path)
print(f"Successfully generated updated {excel_path} with {len(items)} rows and new CV Gap analysis columns!")
