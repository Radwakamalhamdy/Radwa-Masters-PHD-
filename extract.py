import sys
import io
import json
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

items = []

# File 1:
try:
    wb1 = openpyxl.load_workbook('Masters_PhD_Scholarships_Radwa_Oct2026.xlsx', data_only=True)
    ws1 = wb1['Master-PhD Scholarships']
    rows1 = list(ws1.iter_rows(values_only=True))
    for r in rows1[1:]:
        if r[0]:
            items.append({
                'source': 'Radwa_Oct2026',
                'name': str(r[0]).strip(),
                'link': str(r[1]).strip() if r[1] else '',
                'requirements': str(r[2]).strip() if r[2] else '',
                'start_date': str(r[3]).strip() if r[3] else '',
                'deadline': str(r[4]).strip() if r[4] else '',
                'notes': str(r[5]).strip() if r[5] else ''
            })
except Exception as e:
    print(f'Error reading file 1: {e}')

# File 2:
try:
    wb2 = openpyxl.load_workbook('scholarships_master_phd.xlsx', data_only=True)
    ws2 = wb2['المنح']
    rows2 = list(ws2.iter_rows(values_only=True))
    for r in rows2[1:]:
        if len(r) > 1 and r[1]:
            items.append({
                'source': 'scholarships_initial',
                'name': str(r[1]).strip(),
                'link': str(r[2]).strip() if r[2] else '',
                'requirements': str(r[3]).strip() if r[3] else '',
                'start_date': str(r[4]).strip() if r[4] else '',
                'deadline': str(r[5]).strip() if r[5] else '',
                'notes': ''
            })
except Exception as e:
    print(f'Error reading file 2: {e}')

print(f'Total extracted entries: {len(items)}')
with open('extracted_raw.json', 'w', encoding='utf-8') as f:
    json.dump(items, f, ensure_ascii=False, indent=2)
print('Saved to extracted_raw.json successfully!')
